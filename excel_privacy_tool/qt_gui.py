from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from PySide6.QtCore import Qt, Signal
from PySide6.QtGui import QFont
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QFormLayout,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMessageBox,
    QPushButton,
    QSpinBox,
    QStatusBar,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from .core import (
    APP_NAME,
    APP_VERSION,
    PersonColumns,
    PersonConflictError,
    PrivacyToolError,
    anonymize_workbook_multi,
    normalize_ranges,
    mapping_is_encrypted,
    restore_html,
)


class DropLineEdit(QLineEdit):
    fileDropped = Signal(str)

    def __init__(self, accept_files: bool, parent=None):
        super().__init__(parent)
        self.accept_files = accept_files
        self.setAcceptDrops(accept_files)
        if accept_files:
            self.setPlaceholderText("参照ボタン、またはファイルをドラッグ＆ドロップ")

    def dragEnterEvent(self, event):
        if self.accept_files and any(url.isLocalFile() for url in event.mimeData().urls()):
            event.acceptProposedAction()
            self.setStyleSheet("QLineEdit { border: 2px solid #2878bd; background: #edf5fb; }")
        else:
            event.ignore()

    def dragLeaveEvent(self, event):
        self.setStyleSheet("")
        super().dragLeaveEvent(event)

    def dropEvent(self, event):
        self.setStyleSheet("")
        local_files = [url.toLocalFile() for url in event.mimeData().urls() if url.isLocalFile()]
        if local_files:
            path = local_files[0]
            self.setText(path)
            self.fileDropped.emit(path)
            event.acceptProposedAction()
        else:
            event.ignore()


class PathRow(QWidget):
    def __init__(self, file_filter: str, save: bool = False, parent=None):
        super().__init__(parent)
        self.file_filter = file_filter
        self.save = save
        self.edit = DropLineEdit(accept_files=not save)
        self.edit.fileDropped.connect(self._file_dropped)
        button = QPushButton("参照...")
        button.clicked.connect(self.choose)
        layout = QHBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(self.edit, 1)
        layout.addWidget(button)

    def choose(self):
        if self.save:
            path, _ = QFileDialog.getSaveFileName(self, "保存先を選択", self.edit.text(), self.file_filter)
        else:
            path, _ = QFileDialog.getOpenFileName(self, "ファイルを選択", self.edit.text(), self.file_filter)
        if path:
            self.edit.setText(path)
            self.on_chosen(path)

    def on_chosen(self, _path: str):
        pass

    def _file_dropped(self, path: str):
        self.on_chosen(path)

    def text(self) -> str:
        return self.edit.text().strip()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.preview_workbook = None
        self.setWindowTitle(f"{APP_NAME}  {APP_VERSION}")
        screen = QApplication.primaryScreen()
        if screen:
            available = screen.availableGeometry()
            self.resize(min(1120, int(available.width() * 0.9)), min(950, int(available.height() * 0.9)))
        else:
            self.resize(1120, 900)
        self.setMinimumSize(820, 650)
        self._build_ui()

    def _build_ui(self):
        root = QWidget()
        outer = QVBoxLayout(root)
        outer.setContentsMargins(18, 14, 18, 10)
        title = QLabel(APP_NAME)
        title_font = "Hiragino Sans" if sys.platform == "darwin" else "Yu Gothic UI"
        title.setFont(QFont(title_font, 18, QFont.Weight.Bold))
        outer.addWidget(title)
        outer.addWidget(QLabel("個人情報を仮IDに置き換え、AI分析後のHTMLだけを安全に復元します。処理は端末内で完結します。"))
        tabs = QTabWidget()
        tabs.addTab(self._anonymize_tab(), "1. Excelを匿名化")
        tabs.addTab(self._restore_tab(), "2. HTMLを復元")
        outer.addWidget(tabs, 1)
        self.setCentralWidget(root)
        self.setStatusBar(QStatusBar())
        self.statusBar().showMessage("準備完了")

    def _anonymize_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)

        input_box = QGroupBox("入力")
        input_form = QFormLayout(input_box)
        self.excel_input = PathRow("Excelブック (*.xlsx)")
        self.excel_input.on_chosen = self._excel_chosen
        input_form.addRow("Excelファイル", self.excel_input)
        self.sheet_combo = QComboBox()
        self.sheet_combo.currentTextChanged.connect(self._refresh_preview)
        input_form.addRow("シート（列・範囲を追加する対象）", self.sheet_combo)
        layout.addWidget(input_box)

        person_box = QGroupBox("① 人物の列（氏名と社員番号を、同じ人なら1つのIDにまとめる）")
        person_layout = QGridLayout(person_box)
        self.person_name_col = QLineEdit()
        self.person_name_col.setPlaceholderText("例: B")
        self.person_name_col.setMaximumWidth(80)
        self.person_id_col = QLineEdit()
        self.person_id_col.setPlaceholderText("例: A")
        self.person_id_col.setMaximumWidth(80)
        self.person_first_row = QSpinBox()
        self.person_first_row.setRange(1, 1048576)
        self.person_first_row.setValue(2)
        add_person_button = QPushButton("このシートの人物列を追加")
        add_person_button.clicked.connect(self._add_person)
        person_layout.addWidget(QLabel("氏名の列"), 0, 0)
        person_layout.addWidget(self.person_name_col, 0, 1)
        person_layout.addWidget(QLabel("社員番号の列"), 0, 2)
        person_layout.addWidget(self.person_id_col, 0, 3)
        person_layout.addWidget(QLabel("データ開始行"), 0, 4)
        person_layout.addWidget(self.person_first_row, 0, 5)
        person_layout.addWidget(add_person_button, 0, 6)
        person_layout.setColumnStretch(7, 1)
        person_hint = QLabel(
            "シートごとに列が違う場合は、シートを切り替えて列を指定します（片方の列だけでも可）。"
            "同じ行の氏名と社員番号は同一人物として扱います。見出し行は開始行で除外します。"
        )
        person_hint.setWordWrap(True)
        person_hint.setStyleSheet("color:#555")
        person_layout.addWidget(person_hint, 1, 0, 1, 8)
        layout.addWidget(person_box)

        range_box = QGroupBox("② その他の個人情報のセル範囲（住所・電話番号など。値ごとに別IDになります）")
        range_layout = QGridLayout(range_box)
        self.range_edit = QLineEdit()
        self.range_edit.setPlaceholderText("例: B2:B100 または B2:D100")
        self.range_edit.returnPressed.connect(self._add_range)
        add_button = QPushButton("追加")
        add_button.clicked.connect(self._add_range)
        remove_button = QPushButton("選択を削除")
        remove_button.clicked.connect(self._remove_ranges)
        preview_button = QPushButton("プレビュー選択を追加")
        preview_button.clicked.connect(self._add_preview_selection)
        self.range_list = QListWidget()
        self.range_list.setMinimumHeight(90)
        self.range_list.setMaximumHeight(140)
        range_layout.addWidget(QLabel("セル範囲"), 0, 0)
        range_layout.addWidget(QLabel("指定済み"), 2, 0)
        range_layout.addWidget(self.range_edit, 0, 1)
        range_layout.addWidget(add_button, 0, 2)
        range_layout.addWidget(remove_button, 0, 3)
        range_layout.addWidget(preview_button, 0, 4)
        hint = QLabel(
            "シートを切り替えて、シートごとに追加できます。同じ値は別のシートでも同じIDになります。"
            "列見出しなど、個人情報ではないセルは除外してください。"
        )
        hint.setWordWrap(True)
        hint.setStyleSheet("color:#555")
        range_layout.addWidget(hint, 1, 1, 1, 4)
        range_layout.addWidget(self.range_list, 2, 1, 1, 4)
        layout.addWidget(range_box)

        layout.addWidget(QLabel("Excelプレビュー（先頭200行・50列）"))
        self.preview = QTableWidget()
        self.preview.setMinimumHeight(240)
        self.preview.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.preview.setAlternatingRowColors(True)
        layout.addWidget(self.preview, 1)

        output_box = QGroupBox("保存先と対応表")
        output_form = QFormLayout(output_box)
        self.excel_output = PathRow("Excelブック (*.xlsx)", save=True)
        self.map_output = PathRow("対応表 (*.epbmap)", save=True)
        output_form.addRow("匿名化Excel", self.excel_output)
        output_form.addRow("対応表", self.map_output)
        self.encrypt_check = QCheckBox("対応表をパスワードで暗号化する（推奨）")
        self.encrypt_check.setChecked(True)
        self.encrypt_check.toggled.connect(self._encryption_toggled)
        output_form.addRow("", self.encrypt_check)
        self.password = QLineEdit()
        self.password.setEchoMode(QLineEdit.EchoMode.Password)
        self.password_confirm = QLineEdit()
        self.password_confirm.setEchoMode(QLineEdit.EchoMode.Password)
        output_form.addRow("パスワード", self.password)
        output_form.addRow("確認", self.password_confirm)
        action_row = QWidget()
        action_layout = QHBoxLayout(action_row)
        action_layout.setContentsMargins(0, 3, 0, 0)
        action_layout.addStretch()
        self.anonymize_button = QPushButton("匿名化して保存")
        self.anonymize_button.setMinimumHeight(34)
        self.anonymize_button.clicked.connect(self._anonymize)
        action_layout.addWidget(self.anonymize_button)
        output_form.addRow("", action_row)
        layout.addWidget(output_box)
        return tab

    def _restore_tab(self):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        box = QGroupBox("AI分析後のHTMLを復元")
        form = QFormLayout(box)
        self.html_input = PathRow("HTMLファイル (*.html *.htm)")
        self.html_input.on_chosen = self._html_chosen
        self.restore_map = PathRow("対応表 (*.epbmap)")
        self.restore_map.on_chosen = self._restore_map_chosen
        self.restore_password = QLineEdit()
        self.restore_password.setEchoMode(QLineEdit.EchoMode.Password)
        self.html_output = PathRow("HTMLファイル (*.html)", save=True)
        form.addRow("HTMLファイル", self.html_input)
        form.addRow("対応表", self.restore_map)
        form.addRow("パスワード", self.restore_password)
        password_hint = QLabel("暗号化していない対応表では空欄のままにします。")
        password_hint.setStyleSheet("color:#555")
        form.addRow("", password_hint)
        form.addRow("復元HTML", self.html_output)
        action_row = QWidget()
        action_layout = QHBoxLayout(action_row)
        action_layout.setContentsMargins(0, 3, 0, 0)
        action_layout.addStretch()
        self.restore_button = QPushButton("IDを元の値に戻して保存")
        self.restore_button.setMinimumHeight(34)
        self.restore_button.clicked.connect(self._restore)
        action_layout.addWidget(self.restore_button)
        form.addRow("", action_row)
        layout.addWidget(box)

        note = QGroupBox("安全上の注意")
        note_layout = QVBoxLayout(note)
        note_layout.addWidget(QLabel(
            "・AIへ渡すのは匿名化済みExcelだけにし、対応表は渡さないでください。\n"
            "・復元HTMLには個人情報が含まれます。保存先と共有範囲に注意してください。\n"
            "・AIが仮IDを書き換えた場合は復元できません。処理後の検査結果を確認してください。"
        ))
        layout.addWidget(note)
        layout.addStretch()
        return tab

    def _excel_chosen(self, path: str):
        source = Path(path)
        self.excel_output.edit.setText(str(source.with_name(f"{source.stem}_anonymized.xlsx")))
        self.map_output.edit.setText(str(source.with_name(f"{source.stem}_anonymized.epbmap")))
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        try:
            if self.preview_workbook:
                self.preview_workbook.close()
            self.preview_workbook = load_workbook(path, read_only=True, data_only=False)
            self.range_list.clear()
            self.sheet_combo.clear()
            self.sheet_combo.addItems(self.preview_workbook.sheetnames)
            self._refresh_preview(self.sheet_combo.currentText())
            self.statusBar().showMessage(f"読み込み完了: {source.name}")
        except Exception as exc:
            self._error("Excel読み込みエラー", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def _refresh_preview(self, sheet_name: str):
        if not self.preview_workbook or sheet_name not in self.preview_workbook.sheetnames:
            return
        sheet = self.preview_workbook[sheet_name]
        columns = min(max(sheet.max_column, 1), 50)
        rows = min(max(sheet.max_row, 1), 200)
        self.preview.clear()
        self.preview.setColumnCount(columns)
        self.preview.setRowCount(rows)
        self.preview.setHorizontalHeaderLabels([get_column_letter(i) for i in range(1, columns + 1)])
        self.preview.setVerticalHeaderLabels([str(i) for i in range(1, rows + 1)])
        for r, row in enumerate(sheet.iter_rows(min_row=1, max_row=rows, max_col=columns, values_only=True)):
            for c, value in enumerate(row):
                if value is not None:
                    self.preview.setItem(r, c, QTableWidgetItem(str(value)))
        self.preview.resizeColumnsToContents()
        for col in range(columns):
            self.preview.setColumnWidth(col, min(max(self.preview.columnWidth(col), 80), 220))

    def _add_target(self, sheet: str, cell_range: str):
        label = f"{sheet}!{cell_range}"
        if self.range_list.findItems(label, Qt.MatchFlag.MatchExactly):
            return
        item = QListWidgetItem(label)
        item.setData(Qt.ItemDataRole.UserRole, ("range", sheet, cell_range))
        self.range_list.addItem(item)

    def _add_person(self):
        sheet = self.sheet_combo.currentText()
        name_col = self.person_name_col.text().strip().upper()
        id_col = self.person_id_col.text().strip().upper()
        if not sheet:
            self._warning("シート未選択", "先にExcelファイルを選択してください。")
            return
        if not name_col and not id_col:
            self._warning("列未指定", "氏名の列と社員番号の列の、少なくとも一方を入力してください。")
            return
        for label, col in (("氏名", name_col), ("社員番号", id_col)):
            if col:
                try:
                    column_index_from_string(col)
                except ValueError:
                    self._warning("列指定エラー", f"{label}の列は A, B, AA のような列記号で入力してください: {col}")
                    return
        if name_col and name_col == id_col:
            self._warning("列指定エラー", "氏名の列と社員番号の列に同じ列は指定できません。")
            return
        first_row = self.person_first_row.value()
        parts = []
        if name_col:
            parts.append(f"氏名={name_col}")
        if id_col:
            parts.append(f"社員番号={id_col}")
        label = f"{sheet}! 人物列 {' '.join(parts)}（{first_row}行目〜）"
        if self.range_list.findItems(label, Qt.MatchFlag.MatchExactly):
            return
        item = QListWidgetItem(label)
        item.setData(Qt.ItemDataRole.UserRole, ("person", sheet, name_col or None, id_col or None, first_row))
        self.range_list.addItem(item)
        self.person_name_col.clear()
        self.person_id_col.clear()

    def _add_range(self):
        sheet = self.sheet_combo.currentText()
        text = self.range_edit.text().strip()
        if not sheet:
            self._warning("シート未選択", "先にExcelファイルを選択してください。")
            return
        if text:
            try:
                ranges = normalize_ranges([text])
            except PrivacyToolError as exc:
                self._warning("範囲エラー", str(exc))
                return
            for value in ranges:
                self._add_target(sheet, value)
        self.range_edit.clear()

    def _remove_ranges(self):
        for item in self.range_list.selectedItems():
            self.range_list.takeItem(self.range_list.row(item))

    def _add_preview_selection(self):
        selections = self.preview.selectedRanges()
        if not selections:
            self._warning("範囲未選択", "プレビュー上で匿名化するセルを選択してください。")
            return
        for selected in selections:
            start = f"{get_column_letter(selected.leftColumn() + 1)}{selected.topRow() + 1}"
            end = f"{get_column_letter(selected.rightColumn() + 1)}{selected.bottomRow() + 1}"
            value = start if start == end else f"{start}:{end}"
            self._add_target(self.sheet_combo.currentText(), value)

    def _encryption_toggled(self, enabled: bool):
        self.password.setEnabled(enabled)
        self.password_confirm.setEnabled(enabled)
        if not enabled:
            self.password.clear()
            self.password_confirm.clear()

    def _html_chosen(self, path: str):
        source = Path(path)
        self.html_output.edit.setText(str(source.with_name(f"{source.stem}_restored{source.suffix or '.html'}")))

    def _restore_map_chosen(self, path: str):
        if mapping_is_encrypted(path):
            self.statusBar().showMessage("暗号化された対応表です。パスワードを入力してください。")
            self.restore_password.setFocus()
        else:
            self.statusBar().showMessage("暗号化されていない対応表です。")

    def _anonymize(self):
        if self.range_edit.text().strip():
            self._add_range()
        entries = [
            self.range_list.item(i).data(Qt.ItemDataRole.UserRole)
            for i in range(self.range_list.count())
        ]
        targets = [e for e in entries if e[0] == "range"]
        person_columns = [
            PersonColumns(sheet=e[1], name_col=e[2], id_col=e[3], first_row=e[4])
            for e in entries if e[0] == "person"
        ]
        if not all([self.excel_input.text(), self.excel_output.text(), self.map_output.text()]):
            self._warning("入力不足", "入力ファイルと保存先を指定してください。")
            return
        if not targets and not person_columns:
            self._warning("範囲未指定", "人物の列、またはセル範囲を1つ以上追加してください。")
            return
        grouped: dict[str, list[str]] = {}
        for _kind, sheet, cell_range in targets:
            grouped.setdefault(sheet, []).append(cell_range)
        password = None
        if self.encrypt_check.isChecked():
            password = self.password.text()
            if not password:
                self._warning("パスワード未入力", "暗号化パスワードを入力してください。")
                return
            if password != self.password_confirm.text():
                self._warning("パスワード不一致", "確認用パスワードと一致しません。")
                return
        elif QMessageBox.question(
            self, "暗号化なしで保存", "対応表には元の個人情報が含まれます。暗号化せずに保存しますか？",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        ) != QMessageBox.StandardButton.Yes:
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        self.statusBar().showMessage("匿名化して保存中...")
        QApplication.processEvents()
        try:
            def run(allow_conflicts: bool):
                return anonymize_workbook_multi(
                    self.excel_input.text(), self.excel_output.text(), self.map_output.text(),
                    list(grouped.items()), password, person_columns, allow_conflicts
                )

            try:
                result = run(False)
            except PersonConflictError as conflict:
                QApplication.restoreOverrideCursor()
                proceed = self._confirm_conflicts(conflict.conflicts)
                QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
                if not proceed:
                    self.statusBar().showMessage("キャンセルしました（何も保存していません）")
                    return
                result = run(True)
            self.statusBar().showMessage("匿名化が完了しました")
            person_line = f"人物（IDの数）: {result.person_count}\n" if person_columns else ""
            conflict_line = (
                f"\n※ 複数の社員番号/氏名があった人物: {len(result.conflicts)}件（同一人物としてまとめました）\n"
                if result.conflicts else ""
            )
            sheet_lines = "\n".join(
                f"  ・{name}: {count}セル" for name, count in result.replaced_by_sheet.items()
            )
            QMessageBox.information(
                self, "匿名化完了",
                f"匿名化Excelと対応表を保存しました。\n\n置換セル: {result.replaced_cells}\n{sheet_lines}\n{person_line}"
                f"異なる値（シートをまたぐ同一値は1つ）: {result.unique_values}\n空白スキップ: {result.skipped_blank_cells}\n"
                f"数式スキップ: {result.skipped_formula_cells}\n{conflict_line}\nAIへ渡すのは匿名化Excelだけです。"
            )
        except Exception as exc:
            self.statusBar().showMessage("処理に失敗しました")
            self._error("処理エラー", str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def _confirm_conflicts(self, conflicts) -> bool:
        lines = [c.describe() for c in conflicts]
        box = QMessageBox(self)
        box.setIcon(QMessageBox.Icon.Warning)
        box.setWindowTitle("確認が必要です")
        box.setText(
            f"同じ人物に複数の社員番号（または、同じ社員番号に複数の氏名）が{len(conflicts)}件見つかりました。\n\n"
            "契約社員から正社員への登用などで社員番号が変わった場合は、同一人物として1つのIDにまとめます。\n"
            "ただし同姓同名の別人の場合も、同じIDにまとめられてしまいます。\n"
            "「詳細を表示」で内容を確認してください。まとめて続行しますか？"
        )
        box.setDetailedText("\n\n".join(lines))
        proceed = box.addButton("同一人物としてまとめて続行", QMessageBox.ButtonRole.AcceptRole)
        cancel = box.addButton("キャンセル（保存しない）", QMessageBox.ButtonRole.RejectRole)
        box.setDefaultButton(cancel)
        box.exec()
        return box.clickedButton() is proceed

    def _restore(self):
        if not all([self.html_input.text(), self.restore_map.text(), self.html_output.text()]):
            self._warning("入力不足", "HTML、対応表、保存先を指定してください。")
            return
        QApplication.setOverrideCursor(Qt.CursorShape.WaitCursor)
        self.statusBar().showMessage("HTMLを復元中...")
        QApplication.processEvents()
        try:
            result = restore_html(
                self.html_input.text(), self.restore_map.text(), self.html_output.text(),
                self.restore_password.text() or None
            )
            warnings = []
            if result.missing_expected_tokens:
                warnings.append(f"HTML内に見つからなかったID: {len(result.missing_expected_tokens)}件")
            if result.unknown_tokens:
                warnings.append(f"対応表にないID: {len(result.unknown_tokens)}件")
            detail = "\n".join(warnings) if warnings else "ID検査: 問題なし"
            self.statusBar().showMessage("HTMLの復元が完了しました")
            QMessageBox.information(self, "復元完了", f"復元HTMLを保存しました。\n\n置換箇所: {result.replacements}\n{detail}")
        except Exception as exc:
            self.statusBar().showMessage("処理に失敗しました")
            title = "処理エラー" if isinstance(exc, PrivacyToolError) else "予期しないエラー"
            self._error(title, str(exc))
        finally:
            QApplication.restoreOverrideCursor()

    def _warning(self, title: str, message: str):
        QMessageBox.warning(self, title, message)

    def _error(self, title: str, message: str):
        QMessageBox.critical(self, title, message)

    def closeEvent(self, event):
        if self.preview_workbook:
            self.preview_workbook.close()
        event.accept()


def main():
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setApplicationVersion(APP_VERSION)
    app.setOrganizationName("Personal Data Tools")
    app.setStyle("Fusion")
    font_name = "Hiragino Sans" if sys.platform == "darwin" else "Yu Gothic UI"
    app.setFont(QFont(font_name, 10))
    window = MainWindow()
    window.show()
    sys.exit(app.exec())
