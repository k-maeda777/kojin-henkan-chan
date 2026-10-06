from __future__ import annotations

import base64
import html
import json
import os
import re
import secrets
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from cryptography.exceptions import InvalidTag
from cryptography.hazmat.primitives import hashes
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries


APP_NAME = "個人情報変換ちゃん"
APP_VERSION = "0.3.0"
ENCRYPTED_FORMAT = "excel-privacy-bridge-map-v1"
PLAIN_FORMAT = "excel-privacy-bridge-map-plain-v1"
KDF_ITERATIONS = 600_000
TOKEN_RE = re.compile(r"\bPII_[A-Z0-9]{6}_[0-9]{6}(?:_[EN][0-9]+)?\b")


class PrivacyToolError(Exception):
    pass


class WrongPasswordError(PrivacyToolError):
    pass


@dataclass
class AnonymizeResult:
    output_path: str
    map_path: str
    replaced_cells: int
    unique_values: int
    skipped_blank_cells: int
    skipped_formula_cells: int
    replaced_by_sheet: dict[str, int] = field(default_factory=dict)
    person_count: int = 0
    conflicts: list = field(default_factory=list)


@dataclass
class RestoreResult:
    output_path: str
    replacements: int
    missing_expected_tokens: list[str]
    unknown_tokens: list[str]


def _derive_key(password: str, salt: bytes, iterations: int = KDF_ITERATIONS) -> bytes:
    if not password:
        raise PrivacyToolError("暗号化にはパスワードが必要です。")
    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(), length=32, salt=salt, iterations=iterations
    )
    return kdf.derive(password.encode("utf-8"))


def save_mapping(mapping: dict, path: str | Path, password: str | None) -> None:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(mapping, ensure_ascii=False, separators=(",", ":")).encode("utf-8")

    if password:
        salt = os.urandom(16)
        nonce = os.urandom(12)
        key = _derive_key(password, salt)
        encrypted = AESGCM(key).encrypt(nonce, payload, ENCRYPTED_FORMAT.encode("ascii"))
        wrapper = {
            "format": ENCRYPTED_FORMAT,
            "kdf": "PBKDF2-HMAC-SHA256",
            "iterations": KDF_ITERATIONS,
            "salt": base64.b64encode(salt).decode("ascii"),
            "nonce": base64.b64encode(nonce).decode("ascii"),
            "ciphertext": base64.b64encode(encrypted).decode("ascii"),
        }
    else:
        wrapper = {"format": PLAIN_FORMAT, "payload": mapping}

    destination.write_text(json.dumps(wrapper, ensure_ascii=False, indent=2), encoding="utf-8")


def load_mapping(path: str | Path, password: str | None = None) -> dict:
    try:
        wrapper = json.loads(Path(path).read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise PrivacyToolError(f"対応表を読み込めません: {exc}") from exc

    format_name = wrapper.get("format")
    if format_name == PLAIN_FORMAT:
        return wrapper["payload"]
    if format_name != ENCRYPTED_FORMAT:
        raise PrivacyToolError("未対応の対応表形式です。")
    if not password:
        raise WrongPasswordError("この対応表は暗号化されています。パスワードを入力してください。")

    try:
        salt = base64.b64decode(wrapper["salt"])
        nonce = base64.b64decode(wrapper["nonce"])
        ciphertext = base64.b64decode(wrapper["ciphertext"])
        iterations = int(wrapper.get("iterations", KDF_ITERATIONS))
        key = _derive_key(password, salt, iterations)
        payload = AESGCM(key).decrypt(
            nonce, ciphertext, ENCRYPTED_FORMAT.encode("ascii")
        )
        return json.loads(payload.decode("utf-8"))
    except (InvalidTag, ValueError, KeyError, json.JSONDecodeError) as exc:
        raise WrongPasswordError("パスワードが違うか、対応表が破損しています。") from exc


def mapping_is_encrypted(path: str | Path) -> bool:
    try:
        wrapper = json.loads(Path(path).read_text(encoding="utf-8"))
        return wrapper.get("format") == ENCRYPTED_FORMAT
    except (OSError, json.JSONDecodeError):
        return False


def normalize_ranges(ranges: Iterable[str]) -> list[str]:
    normalized: list[str] = []
    for raw in ranges:
        for item in raw.split(","):
            value = item.strip().upper().replace("$", "")
            if not value:
                continue
            try:
                range_boundaries(value)
            except ValueError as exc:
                raise PrivacyToolError(f"セル範囲が正しくありません: {value}") from exc
            normalized.append(value)
    if not normalized:
        raise PrivacyToolError("1つ以上のセル範囲を指定してください。")
    return list(dict.fromkeys(normalized))


@dataclass
class PersonColumns:
    """人物（氏名・社員番号）の列指定。同じ行の氏名と社員番号は「同一人物」として1つのIDにまとめる。"""

    sheet: str
    name_col: str | None = None
    id_col: str | None = None
    first_row: int = 2
    last_row: int | None = None


@dataclass
class PersonConflict:
    kind: str  # "multi_id": 同一人物に複数の社員番号 / "multi_name": 同一社員番号に複数の氏名
    values: list[str]
    context: list[str]
    locations: list[str]

    def describe(self) -> str:
        if self.kind == "multi_id":
            who = "・".join(self.context) if self.context else "（氏名なし）"
            head = f"{who} に複数の社員番号: {' / '.join(self.values)}"
        else:
            head = f"社員番号 {' / '.join(self.context)} に複数の氏名: {' / '.join(self.values)}"
        return f"{head}\n    場所: {', '.join(self.locations)}"


class PersonConflictError(PrivacyToolError):
    def __init__(self, conflicts: list[PersonConflict]):
        self.conflicts = conflicts
        super().__init__(f"同一人物に複数の社員番号（または氏名）が見つかりました: {len(conflicts)}件")


def _identity_key(value: str) -> str:
    """同一人物の判定キー。全角/半角・空白・英字の大小の違いを吸収する。"""
    return "".join(unicodedata.normalize("NFKC", value).split()).casefold()


def _cell_text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excelの数値 1001.0 を "1001" として扱う
    return str(value)


def _column_index(letter: str | None, label: str) -> int | None:
    if letter is None or not str(letter).strip():
        return None
    try:
        return column_index_from_string(str(letter).strip().upper())
    except ValueError as exc:
        raise PrivacyToolError(f"{label}の列指定が正しくありません: {letter}") from exc


def anonymize_workbook(
    input_path: str | Path,
    output_path: str | Path,
    map_path: str | Path,
    sheet_name: str,
    ranges: Iterable[str],
    password: str | None = None,
) -> AnonymizeResult:
    """1シートだけを匿名化する（従来のインターフェース）。"""
    return anonymize_workbook_multi(
        input_path, output_path, map_path, [(sheet_name, ranges)], password
    )


class _Node:
    __slots__ = ("key", "original", "is_number", "locations")

    def __init__(self, key: str, original: str, is_number: bool):
        self.key = key
        self.original = original
        self.is_number = is_number
        self.locations: list[str] = []


def anonymize_workbook_multi(
    input_path: str | Path,
    output_path: str | Path,
    map_path: str | Path,
    targets: Iterable[tuple[str, Iterable[str]]] = (),
    password: str | None = None,
    person_columns: Iterable[PersonColumns] = (),
    allow_conflicts: bool = False,
) -> AnonymizeResult:
    """複数シートをまとめて匿名化する。

    targets: (シート名, セル範囲のリスト)。範囲内の値を、値ごとに仮IDへ置き換える。
    person_columns: 氏名列・社員番号列の指定。同じ行の氏名と社員番号を同一人物として1つのIDにまとめる。
        氏名セルは基本ID（PII_XXXXXX_000001）、社員番号セルは枝番付き（PII_XXXXXX_000001_E1）になる。
    同じ値（空白・全角半角の違いは無視）は、シートをまたいでも同じ人物になる。
    同一人物に複数の社員番号（または同一社員番号に複数の氏名）があると PersonConflictError を送出する。
    allow_conflicts=True なら、同一人物としてまとめて続行する。
    """
    source = Path(input_path)
    destination = Path(output_path)
    mapping_destination = Path(map_path)
    if source.resolve() == destination.resolve():
        raise PrivacyToolError("入力ファイルと出力ファイルは別名にしてください。")

    # 同じシートが複数回指定されたら範囲をまとめる（順序は維持）
    sheet_ranges: dict[str, list[str]] = {}
    for sheet_name, ranges in targets:
        merged = sheet_ranges.setdefault(sheet_name, [])
        merged.extend(normalize_ranges(ranges))
    sheet_ranges = {name: list(dict.fromkeys(r)) for name, r in sheet_ranges.items()}

    persons: list[tuple[PersonColumns, int | None, int | None]] = []
    for spec in person_columns:
        name_idx = _column_index(spec.name_col, "氏名")
        id_idx = _column_index(spec.id_col, "社員番号")
        if name_idx is None and id_idx is None:
            raise PrivacyToolError(f"{spec.sheet}: 氏名列と社員番号列の少なくとも一方を指定してください。")
        if name_idx is not None and name_idx == id_idx:
            raise PrivacyToolError(f"{spec.sheet}: 氏名列と社員番号列に同じ列は指定できません。")
        if spec.first_row < 1:
            raise PrivacyToolError(f"{spec.sheet}: 開始行は1以上にしてください。")
        persons.append((spec, name_idx, id_idx))

    if not sheet_ranges and not persons:
        raise PrivacyToolError("1つ以上のシートと範囲、または人物の列を指定してください。")

    try:
        workbook = load_workbook(source)
    except Exception as exc:
        raise PrivacyToolError(f"Excelファイルを開けません: {exc}") from exc
    for sheet_name in [*sheet_ranges, *(spec.sheet for spec, _, _ in persons)]:
        if sheet_name not in workbook.sheetnames:
            raise PrivacyToolError(f"シートが見つかりません: {sheet_name}")

    nodes: dict[str, _Node] = {}
    parent: dict[str, str] = {}
    cell_refs: list[tuple[str, str, str]] = []  # (sheet, coordinate, node key)
    processed: dict[str, set[str]] = {}
    counters = {"blank": 0, "formula": 0}

    def find(key: str) -> str:
        while parent[key] != key:
            parent[key] = parent[parent[key]]
            key = parent[key]
        return key

    def union(a: str, b: str) -> None:
        ra, rb = find(a), find(b)
        if ra != rb:
            parent[rb] = ra

    def visit(sheet_name: str, cell, is_number: bool) -> str | None:
        done = processed.setdefault(sheet_name, set())
        if cell.coordinate in done:
            return None
        done.add(cell.coordinate)
        value = cell.value
        if value is None or (isinstance(value, str) and not value.strip()):
            counters["blank"] += 1
            return None
        if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
            counters["formula"] += 1
            return None
        original = _cell_text(value)
        key = _identity_key(original)
        node = nodes.get(key)
        if node is None:
            node = nodes[key] = _Node(key, original, is_number)
            parent[key] = key
        elif not is_number:
            node.is_number = False
        node.locations.append(f"{sheet_name}!{cell.coordinate}")
        cell_refs.append((sheet_name, cell.coordinate, key))
        return key

    # 人物列を先に処理する（同じセルが一般範囲にも含まれても、人物としての役割を優先する）
    for spec, name_idx, id_idx in persons:
        sheet = workbook[spec.sheet]
        last_row = spec.last_row or sheet.max_row
        for row in range(spec.first_row, last_row + 1):
            name_key = id_key = None
            if name_idx is not None:
                name_key = visit(spec.sheet, sheet.cell(row=row, column=name_idx), False)
            if id_idx is not None:
                id_key = visit(spec.sheet, sheet.cell(row=row, column=id_idx), True)
            if name_key and id_key:
                union(name_key, id_key)

    for sheet_name, normalized_ranges in sheet_ranges.items():
        sheet = workbook[sheet_name]
        for cell_range in normalized_ranges:
            min_col, min_row, max_col, max_row = range_boundaries(cell_range)
            min_col = min_col or 1
            min_row = min_row or 1
            max_col = max_col or sheet.max_column
            max_row = max_row or sheet.max_row
            for row in sheet.iter_rows(
                min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col
            ):
                for cell in row:
                    visit(sheet_name, cell, False)

    if not cell_refs:
        raise PrivacyToolError("置換できるセルがありませんでした。範囲を確認してください。")

    # 人物（連結成分）ごとにまとめる。出現順を保つ
    groups: dict[str, list[_Node]] = {}
    for key, node in nodes.items():
        groups.setdefault(find(key), []).append(node)

    conflicts: list[PersonConflict] = []
    for members in groups.values():
        numbers = [n for n in members if n.is_number]
        names = [n for n in members if not n.is_number]
        if len(numbers) >= 2:
            conflicts.append(PersonConflict(
                "multi_id", [n.original for n in numbers], [n.original for n in names],
                [loc for n in members for loc in n.locations[:2]],
            ))
        if len(names) >= 2 and numbers:
            conflicts.append(PersonConflict(
                "multi_name", [n.original for n in names], [n.original for n in numbers],
                [loc for n in members for loc in n.locations[:2]],
            ))
    if conflicts and not allow_conflicts:
        raise PersonConflictError(conflicts)

    prefix = secrets.token_hex(3).upper()
    node_token: dict[str, str] = {}
    token_to_original: dict[str, str] = {}
    for index, members in enumerate(groups.values(), start=1):
        base = f"PII_{prefix}_{index:06d}"
        names = [n for n in members if not n.is_number]
        numbers = [n for n in members if n.is_number]
        primary = names[0] if names else numbers[0]
        node_token[primary.key] = base
        token_to_original[base] = primary.original
        for position, node in enumerate(names[1:], start=2):
            node_token[node.key] = f"{base}_N{position}"
            token_to_original[node_token[node.key]] = node.original
        for position, node in enumerate([n for n in numbers if n is not primary], start=1):
            node_token[node.key] = f"{base}_E{position}"
            token_to_original[node_token[node.key]] = node.original

    cell_records: list[dict[str, str]] = []
    replaced_by_sheet: dict[str, int] = {}
    for sheet_name, coordinate, key in cell_refs:
        token = node_token[key]
        workbook[sheet_name][coordinate].value = token
        cell_records.append({"sheet": sheet_name, "cell": coordinate, "token": token})
        replaced_by_sheet[sheet_name] = replaced_by_sheet.get(sheet_name, 0) + 1

    first_sheet = next(iter(sheet_ranges), None) or persons[0][0].sheet
    mapping = {
        "schema": 1,
        "app": APP_NAME,
        "app_version": APP_VERSION,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "source_filename": source.name,
        "anonymized_filename": destination.name,
        "sheet": first_sheet,
        "ranges": sheet_ranges.get(first_sheet, []),
        "targets": [{"sheet": name, "ranges": r} for name, r in sheet_ranges.items()],
        "person_columns": [
            {
                "sheet": spec.sheet, "name_col": spec.name_col, "id_col": spec.id_col,
                "first_row": spec.first_row, "last_row": spec.last_row,
            }
            for spec, _, _ in persons
        ],
        "token_prefix": f"PII_{prefix}_",
        "values": token_to_original,
        "cells": cell_records,
    }

    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        save_mapping(mapping, mapping_destination, password)
    except Exception as exc:
        if destination.exists():
            destination.unlink(missing_ok=True)
        raise PrivacyToolError(f"ファイルを保存できません: {exc}") from exc

    return AnonymizeResult(
        output_path=str(destination),
        map_path=str(mapping_destination),
        replaced_cells=len(cell_refs),
        unique_values=len(token_to_original),
        skipped_blank_cells=counters["blank"],
        skipped_formula_cells=counters["formula"],
        replaced_by_sheet=replaced_by_sheet,
        person_count=len(groups),
        conflicts=conflicts,
    )


def _decode_html(raw: bytes) -> tuple[str, str, bytes]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8"), "utf-8", b"\xef\xbb\xbf"
    if raw.startswith(b"\xff\xfe"):
        return raw[2:].decode("utf-16-le"), "utf-16-le", b"\xff\xfe"
    if raw.startswith(b"\xfe\xff"):
        return raw[2:].decode("utf-16-be"), "utf-16-be", b"\xfe\xff"
    try:
        return raw.decode("utf-8"), "utf-8", b""
    except UnicodeDecodeError:
        try:
            return raw.decode("cp932"), "cp932", b""
        except UnicodeDecodeError as exc:
            raise PrivacyToolError("HTMLの文字コードを判定できません。UTF-8で保存し直してください。") from exc


def restore_html(
    html_path: str | Path,
    map_path: str | Path,
    output_path: str | Path,
    password: str | None = None,
) -> RestoreResult:
    try:
        text, encoding, bom = _decode_html(Path(html_path).read_bytes())
    except OSError as exc:
        raise PrivacyToolError(f"HTMLを読み込めません: {exc}") from exc
    mapping = load_mapping(map_path, password)
    values: dict[str, str] = mapping.get("values", {})
    if not values:
        raise PrivacyToolError("対応表に復元データがありません。")

    present_before = {token for token in values if token in text}
    replacements = 0
    for token in sorted(values, key=len, reverse=True):
        count = text.count(token)
        if count:
            # HTML本文で安全に表示できる形へ戻す。通常の日本語氏名は変化しない。
            text = text.replace(token, html.escape(str(values[token]), quote=True))
            replacements += count

    expected = set(values)
    missing = sorted(expected - present_before)
    unknown = sorted(set(TOKEN_RE.findall(text)) - expected)
    destination = Path(output_path)
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(bom + text.encode(encoding))
    except OSError as exc:
        raise PrivacyToolError(f"復元HTMLを保存できません: {exc}") from exc

    return RestoreResult(
        output_path=str(destination),
        replacements=replacements,
        missing_expected_tokens=missing,
        unknown_tokens=unknown,
    )
